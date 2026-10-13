"""Converts a StoreHippo product export (.xlsx, one sheet per category) into catalog.json for prisma/import-catalog.ts.

    python -I prisma/catalog/storehippo-to-catalog.py <export.xlsx> <catalog.json>

Rules (change them here, not in the data):
  - one product per row; price = selling price, compare_price = original price (shown struck through)
  - status ACTIVE only when the row is approved AND published; everything else is saved as DRAFT
  - condition NEW unless the name or collection says refurbished/renewed
  - HTML descriptions become plain text; images are upgraded to https
"""
import html, json, re, sys, zipfile, io, collections
from openpyxl import load_workbook

BRANDS = {  # lower-cased source value -> display name (the export spells the same brand several ways)
    "boat": "boAt", "oneplus": "OnePlus", "vivo": "vivo", "lapcare": "Lapcare", "realme": "realme", "samsung": "Samsung",
    "zebronics": "Zebronics", "apple": "Apple", "ooge": "Ooge", "nothing": "Nothing", "cmf by nothing": "CMF by Nothing",
    "poco": "POCO", "celljoy": "Celljoy", "infinix": "Infinix", "lava agni": "Lava", "xiaomi": "Xiaomi", "xiaomi mi": "Xiaomi",
    "lyne": "Lyne", "portronics": "Portronics", "hrx": "HRX", "oppo": "OPPO", "oppo/realme": "OPPO", "motorola": "Motorola",
    "google": "Google", "non branded": "Unbranded",
}
COUNTRY = {"IN": "India", "CN": "China"}


def load_rows(path):
    # StoreHippo exports carry data-validation XML that openpyxl rejects; it is not needed, so drop it from a copy in memory.
    raw = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in raw.infolist():
            data = raw.read(item.filename)
            if re.match(r"xl/worksheets/sheet\d+\.xml$", item.filename):
                data = re.sub(r"<dataValidations\b.*?</dataValidations>", "", data.decode("utf-8"), flags=re.S).encode("utf-8")
            out.writestr(item, data)
    buf.seek(0)
    wb = load_workbook(buf, data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        for r in it:
            d = {header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header) and header[i]}
            if d.get("name *"):
                rows.append(d)
    return rows


def plain_text(src):
    t = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", src or "")
    def table_row(m):
        cells = [re.sub(r"<[^>]+>", " ", c) for c in re.findall(r"(?is)<t[dh]\b[^>]*>(.*?)</t[dh]>", m.group(0))]
        cells = [re.sub(r"\s+", " ", html.unescape(c)).strip() for c in cells]
        cells = [c for c in cells if c]
        return "\n" + (": ".join(cells) if len(cells) == 2 else " · ".join(cells)) + "\n"
    t = re.sub(r"(?is)<tr\b.*?</tr>", table_row, t)
    t = re.sub(r"(?i)<br\s*/?>", "\n", t)
    t = re.sub(r"(?i)<li\b[^>]*>", "\n• ", t)
    t = re.sub(r"(?i)</(p|div|ul|ol|h[1-6]|table|li)>", "\n", t)
    t = re.sub(r"<[^>]+>", "", t)
    t = html.unescape(t).replace("\xa0", " ").replace("�", "•")  # one export row has bullets that were corrupted in transit
    t = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in t.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def months(text):
    m = re.search(r"(\d+)\s*(year|yr|yaer|month)", text or "", re.I)
    if not m:
        return 0
    n = int(m.group(1))
    return n * 12 if m.group(2).lower() != "month" else n


def slugify(s, limit=90):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    if len(s) > limit:
        s = s[:limit].rsplit("-", 1)[0] if "-" in s[:limit] else s[:limit]
    return s.strip("-")


def main(src, dst):
    rows = load_rows(src)
    cats, brands, colls = collections.OrderedDict(), collections.OrderedDict(), collections.OrderedDict()
    products, slugs = [], collections.Counter()
    for r in rows:
        brand_raw = (r.get("brand") or r.get("attributes.Brand *") or r.get("brand_name") or "").strip()
        brand = BRANDS.get(brand_raw.lower(), brand_raw.title() if brand_raw else "Unbranded")
        brand_slug = slugify(brand)
        brands.setdefault(brand_slug, brand)

        cat_slug = slugify(r.get("category *") or r["categories *"].split(":")[-1])
        cat_name = r["categories *"].split(",")[-1].split(":")[-1].strip() or cat_slug.replace("-", " ").title()
        cats.setdefault(cat_slug, cat_name)

        coll_slugs = [slugify(c) for c in re.split(r"[,;]", r.get("collections", "")) if c.strip()]
        for c in coll_slugs:
            colls.setdefault(c, c.replace("-", " ").title())

        price, compare = round(float(r["price *"]) * 100), round(float(r["compare_price *"]) * 100)
        base, sale = max(price, compare), (price if price < max(price, compare) else None)

        text = (r["name *"] + " " + r.get("collections", "")).lower()
        condition = "REFURBISHED" if re.search(r"refurb|refub|refit|renewed", text) else "NEW"  # incl. the export's typo "Refubreshed" and the refurbisher "Refit"
        approved = r.get("approve", "").lower() == "approved" and r.get("publish") == "1"

        slug = slugify(r.get("alias") or r["name *"])
        slugs[slug] += 1
        if slugs[slug] > 1:
            slug = f"{slug}-{slugs[slug]}"

        specs = {"Brand": brand}
        if r.get("model_number"): specs["Model"] = r["model_number"]
        if r.get("option1_name") and r.get("option1_value") and r["option1_name"].lower().startswith("colo"): specs["Colour"] = r["option1_value"].title()
        if r.get("warranty"): specs["Warranty"] = re.sub(r"\s+", " ", r["warranty"]).strip()
        if r.get("country_of_origin *"): specs["Country of origin"] = COUNTRY.get(r["country_of_origin *"].upper(), r["country_of_origin *"])

        images = []
        for i in range(1, 20):
            u = r.get("image1 *" if i == 1 else f"image{i}", "")
            if u:
                u = re.sub(r"^http://", "https://", u)
                if u not in images:
                    images.append(u)

        tags = {brand.lower(), cat_name.lower()}
        if r.get("common_name *"): tags.add(r["common_name *"].lower())
        for t in re.split(r"[,;]", r.get("product_tag", "")):
            if t.strip(): tags.add(t.strip().lower())

        colour = specs.get("Colour")
        products.append({
            "name": re.sub(r"\s+", " ", r["name *"]).strip(), "slug": slug, "sku": r["sku"],
            "description": plain_text(r.get("description *", "")) or r["name *"],
            "condition": condition, "status": "ACTIVE" if approved else "DRAFT",
            "basePrice": base, "salePrice": sale, "categorySlug": cat_slug, "brandSlug": brand_slug,
            "stock": max(0, int(float(r.get("inventory_quantity *") or 0))),
            "variantTitle": colour or "Standard", "warrantyMonths": months(r.get("warranty", "")),
            "specifications": specs, "tags": sorted(tags),
            "seoTitle": r.get("seo.title") or None, "seoDescription": r.get("seo.description") or None,
            "collectionSlugs": coll_slugs, "images": images,
            "sourceStatus": {"approve": r.get("approve", ""), "publish": r.get("publish", "")},
        })

    catalog = {
        "parentCategory": {"slug": "electronics", "name": "Electronics"},
        "categories": [{"slug": s, "name": n, "position": i} for i, (s, n) in enumerate(cats.items())],
        "brands": [{"slug": s, "name": n} for s, n in sorted(brands.items(), key=lambda kv: kv[1].lower())],
        "collections": [{"slug": s, "name": n} for s, n in colls.items()],
        "products": products,
    }
    json.dump(catalog, open(dst, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    by_status = collections.Counter(p["status"] for p in products)
    print(f"products: {len(products)} {dict(by_status)} | categories: {len(cats)} | brands: {len(brands)} | collections: {len(colls)}")
    print("refurbished:", sum(1 for p in products if p["condition"] == "REFURBISHED"), "| zero stock:", sum(1 for p in products if p["stock"] == 0), "| no images:", sum(1 for p in products if not p["images"]))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
