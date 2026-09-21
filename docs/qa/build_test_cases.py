"""Builds docs/qa/manual-test-cases.csv (import it into Google Sheets). Run: python docs/qa/build_test_cases.py"""
import csv, os

LIVE = "the live site"
rows = []
counters = {}

def T(prefix, module, prio, title, pre, steps, expected, device="Desktop Chrome"):
    counters[prefix] = counters.get(prefix, 0) + 1
    steps = "\n".join(f"{i}. {s}" for i, s in enumerate(steps.split(" | "), 1))
    rows.append([f"{prefix}-{counters[prefix]:03d}", module, prio, title, pre, steps, expected, device, "Not run", "", "", "", ""])

# ---------------------------------------------------------------- 1 environment
M = "01 Deployment & environment"; P = "ENV"
T(P, M, "P0", "Live site opens over HTTPS", "None", "Open the site address in a browser | Check the padlock in the address bar", "Home page loads with no security warning; http:// redirects to https://")
T(P, M, "P0", "Backend (Render) is healthy", "None", "Open <render-url>/api/health", "Returns {\"status\":\"ok\"} within a few seconds")
T(P, M, "P0", "Render address does not show the store", "None", "Open <render-url>/ and <render-url>/products", "Redirects to the public site (only /api/* stays on Render)")
T(P, M, "P0", "Private endpoints are not public", "None", "Send POST to <site>/api/internal/db and <render-url>/api/internal/db without the secret", "Both return 404; nothing is exposed")
T(P, M, "P1", "First request after idle (Render cold start)", "Backend idle 20+ min with keep-alive job paused", "Open the site | Add a product to the cart", "Page loads within ~60s, no error page; later requests are fast")
T(P, M, "P1", "Keep-alive cron job runs", "cron-job.org account", "Open the Keep alive job history | Check the last 5 runs", "All runs 200 OK, every 10 min")
T(P, M, "P1", "Release-reservations cron job", "cron-job.org", "Open the job history", "Runs every 10 min, 200 OK (401 means the Bearer secret is wrong)")
T(P, M, "P1", "Reconcile-payments cron job", "cron-job.org", "Open the job history", "Runs every 10 min, 200 OK")
T(P, M, "P1", "Notification dispatch cron job", "cron-job.org", "Open the job history", "Runs every 5 min, 200 OK")
T(P, M, "P0", "Razorpay webhook configured", "Razorpay dashboard access", "Open Webhooks | Check URL and events", "URL = <render-url>/api/payments/webhook/razorpay; only payment.captured and payment.failed selected")
T(P, M, "P1", "Environment variables present", "Vercel + Render dashboards", "Compare with docs/DEPLOYMENT.md table", "Vercel: BACKEND_URL, INTERNAL_API_SECRET, NEXT_PUBLIC_APP_URL. Render: DATABASE_URL, AUTH_SECRET, CRON_SECRET, RAZORPAY_*, INTERNAL_API_SECRET, API_ONLY=true")
T(P, M, "P1", "Custom domain (when connected)", "mobilenmore.online pointed to Vercel", "Open https://mobilenmore.online and https://www.mobilenmore.online", "Both load the site with valid HTTPS; www redirects to the main address; NEXT_PUBLIC_APP_URL updated on Vercel and Render")
T(P, M, "P2", "Database backups", "Supabase dashboard", "Check the backups / point-in-time recovery page", "Backups are enabled; owner knows how to restore")

# ---------------------------------------------------------------- 2 header / nav / search
M = "02 Header, categories & search"; P = "NAV"
T(P, M, "P0", "Header shows on every page", "None", "Open Home, Products, a product, Cart, Login, a policy page", "Announcement bar, logo, search, Shop all, Home, Sign in/Account, Cart and the category strip are present on all")
T(P, M, "P1", "Logo goes home", "On any inner page", "Click the Mobile & More logo", "Goes to the home page")
T(P, M, "P1", "Home button", "On any inner page", "Click Home", "Goes to the home page")
T(P, M, "P0", "Search returns results", "Products exist", "Type 'phone' in the search box | Press Enter", "Listing page with matching products; heading shows the search")
T(P, M, "P1", "Search suggestions", "None", "Type 2-3 letters slowly", "Suggestions dropdown shows matching products/categories; click one opens it")
T(P, M, "P1", "Search with no result", "None", "Search for 'zzzzqqq'", "Friendly 'no products' message, no error")
T(P, M, "P1", "Search misspelling is not shared between visitors", "Two different browsers/devices", "On device A search 'moooobile' | On device B open the search box / suggestions", "Device B does NOT see device A's search as a popular/suggested term")
T(P, M, "P2", "Search special characters", "None", "Search for <script>alert(1)</script> and for % and '", "No script runs, no error; results page renders safely")
T(P, M, "P1", "Shop all button", "None", "Click Shop all", "Opens the full product listing")
T(P, M, "P0", "Category strip links", "None", "Click each category in the strip", "Each opens its category listing; the clicked category is highlighted (and 'Refurbished deals' is NOT highlighted on a normal category)")
T(P, M, "P1", "Category strip highlight on Refurbished deals", "None", "Open the Refurbished deals link", "Only Refurbished deals is highlighted")
T(P, M, "P1", "All categories menu (desktop)", "None", "Click All categories", "Menu lists every category with product counts and a 'Browse the full catalog' link; closes on outside click / Esc")
T(P, M, "P0", "All categories menu (phone) locks background scroll", "Phone", "Open All categories | Try to scroll the page behind", "Menu opens as a bottom sheet; the page behind does not scroll; closing restores scrolling", "Phone Chrome/Safari")
T(P, M, "P1", "Cart badge count", "None", "Add 2 different items to the cart", "Badge shows 2 immediately; updates on remove; no badge when the cart is empty")
T(P, M, "P1", "Account/Sign-in link", "Signed out / signed in", "Look at the header in both states", "Signed out shows Sign in; signed in shows Account (and Wishlist); admin sees an Admin link")
T(P, M, "P2", "Announcement bar scrolls", "None", "Watch the top bar 10 seconds", "Messages scroll smoothly and loop; readable on phone")

# ---------------------------------------------------------------- 3 homepage
M = "03 Home page & content"; P = "HOME"
T(P, M, "P0", "Home page loads with live offers", "Products with sale prices exist", "Open the home page", "Hero shows 'Up to X% off', featured offer cards with price and struck-out old price, no broken images")
T(P, M, "P1", "Hero buttons", "None", "Click Shop all deals and Refurbished", "Each opens the correct listing")
T(P, M, "P1", "Hero numbers are correct", "None", "Compare 'products on offer' and the % with the offers in Admin", "Numbers match the real active offers")
T(P, M, "P1", "Coupon strip", "Active coupons", "Look at the coupon cards | Click a coupon code", "Codes match Admin → Offers; clicking copies the code (confirmation shown); expired/paused coupons are not shown")
T(P, M, "P1", "Shop by budget", "None", "Click each 'Under ₹…' card", "Opens the listing filtered to that price; product counts on the cards are right")
T(P, M, "P1", "Category / brand sections", "None", "Click a category card and a brand", "Correct filtered listing opens")
T(P, M, "P2", "Homepage sections toggle", "Admin access", "In Admin → Homepage switch a section off | Refresh the home page", "Section disappears; switching it on brings it back")
T(P, M, "P1", "Banner shows and can be hidden", "A banner exists", "Admin → Homepage → hide the banner | Refresh the home page | Show it again", "Banner disappears then reappears; no error (no 'invalid cuid')")
T(P, M, "P2", "Animations", "None", "Scroll the home page", "Sections fade/slide in smoothly; no flicker, no layout jump; 'reduce motion' setting disables them")
T(P, M, "P1", "Announcement text edit", "Admin access", "Admin → Homepage → edit top-bar messages, save | Refresh the site", "New messages scroll in the bar; empty list falls back to the defaults")
T(P, M, "P1", "Page density at 100% zoom", "Desktop 1440px+", "Open the site at 100% browser zoom", "Layout looks balanced (same as ~85% before); nothing cut off or overlapping", "Desktop Chrome/Edge/Firefox")
T(P, M, "P1", "Footer sits at the bottom", "Tall window / Mac", "Open the cart when empty, Login, a policy page on a tall window", "Footer touches the bottom of the window; no blank white strip below it", "Mac Chrome/Safari")
T(P, M, "P2", "Footer links", "None", "Click every footer link", "All open the correct page; no 404")

# ---------------------------------------------------------------- 4 listing
M = "04 Product listing & filters"; P = "LST"
T(P, M, "P0", "All products list", "None", "Open Shop all", "Grid of active products with image, name, price, badge; count is right")
T(P, M, "P0", "Product card layout is responsive", "None", "Resize from desktop to phone width", "Cards resize evenly (4→3→2 columns), text not cut off, images not stretched", "Desktop + Phone")
T(P, M, "P1", "Filter by price / condition / brand", "None", "Use each filter, then Apply filters", "Only matching products; Apply filters button has normal width; checkboxes aligned")
T(P, M, "P1", "Filters on phone", "Phone", "Open the filter toggle | Apply", "Filters open as a panel; Apply works; result count updates", "Phone")
T(P, M, "P1", "Sorting", "None", "Sort by price low→high, high→low, newest", "Order is correct each time")
T(P, M, "P2", "Pagination / long lists", "30+ products", "Scroll / go to next page", "All products reachable; no duplicates")
T(P, M, "P1", "Sale price shown", "A product on sale", "Look at its card", "Discount badge, sale price and struck-out original price are correct")
T(P, M, "P1", "Out-of-stock product", "A product with 0 stock", "Open it in the listing", "Shows Sold out; Add to cart disabled")
T(P, M, "P2", "Empty category", "A category with no products", "Open it", "Friendly empty state, not an error")
T(P, M, "P1", "Save/wishlist heart on card", "Signed in", "Click the save button on a card", "Product is added; shows in Account → Wishlist; click again removes")

# ---------------------------------------------------------------- 5 product page
M = "05 Product page"; P = "PDP"
T(P, M, "P0", "Product details render", "None", "Open any product", "Images, name, price, condition/grade, warranty, availability, specifications, what's included, description all show")
T(P, M, "P1", "Image gallery", "Product with several photos", "Click thumbnails", "Main image changes; no broken images")
T(P, M, "P0", "Variant selection and stock", "Product with variants", "Pick each variant | Try quantity above stock", "Price/stock follow the variant; cannot add more than available (clear message)")
T(P, M, "P0", "Add to cart from product page", "None", "Click Add to cart", "Button shows added state; badge increases; item is in the cart")
T(P, M, "P1", "Reviews are visible", "Product with an approved review", "Scroll to Customer reviews", "Approved reviews with stars and name show; average rating at the top matches")
T(P, M, "P0", "Review form only for buyers", "Signed out / signed in without a delivered order", "Open a product and scroll to reviews", "No form; note says only customers who bought and received it can review")
T(P, M, "P0", "Buyer can write a review", "Signed in customer with a DELIVERED, paid order for that product", "Open the product | Submit a rating and 10+ characters", "Success message 'sent for moderation'; not public until approved in Admin → Reviews")
T(P, M, "P1", "Review validation", "Eligible buyer", "Submit under 10 characters", "Clear error; nothing saved")
T(P, M, "P1", "Sensitive data hidden", "None", "View page source / product page", "IMEI or serial numbers are never shown")
T(P, M, "P2", "SEO on product page", "None", "View source", "Unique <title>, description, canonical, product structured data with the right price")
T(P, M, "P2", "Unknown product address", "None", "Open /products/does-not-exist", "Proper 'not found' page")

# ---------------------------------------------------------------- 6 cart
M = "06 Cart"; P = "CRT"
T(P, M, "P0", "Guest can add to cart", "Signed out", "Add a product", "Cart contains it; badge updates; works after a page refresh")
T(P, M, "P0", "Change quantity and remove", "Item in cart", "Click + then - then Remove", "Totals update each time; removing the last item shows the empty cart message")
T(P, M, "P0", "Totals are correct", "2 items", "Compare subtotal, GST, delivery, total with prices", "Math matches (GST inclusive as configured, delivery free over ₹499 or per shipping rules)")
T(P, M, "P0", "Apply a valid coupon", "Active coupon, cart above its minimum", "Enter the code | Apply", "Discount line appears, total drops; coupon box is neat (input and Apply button joined)")
T(P, M, "P1", "Invalid / expired / below-minimum coupon", "None", "Try a wrong code, an expired one, and one needing a higher cart value", "Clear red message for each; total unchanged")
T(P, M, "P1", "Coupon is not case sensitive", "Coupon exists", "Enter it in lower case", "Applies normally")
T(P, M, "P1", "Cart survives sign-in", "Guest with items, existing customer account", "Add items as guest | Log in", "Guest items are merged into the account cart, none lost")
T(P, M, "P1", "Cart survives close/reopen", "Signed in", "Add items | Close the browser | Reopen and sign in", "Cart still has the items")
T(P, M, "P1", "Cart on two devices", "Signed in on phone and laptop", "Add on the phone | Refresh cart on the laptop", "Same cart on both")
T(P, M, "P1", "Stock changes while in cart", "Item in cart", "Admin sets its stock to 0 | Open the cart / checkout", "Clear message that the item is unavailable; cannot pay for it")
T(P, M, "P1", "Empty cart page", "Cart empty", "Open /cart", "Friendly message and a link to shop; footer at the bottom")
T(P, M, "P2", "Very large quantity", "None", "Type 9999 or negative numbers via the URL/API", "Rejected with a message; no crash")

# ---------------------------------------------------------------- 7 auth
M = "07 Sign up, sign in & account access"; P = "AUT"
T(P, M, "P0", "Sign up with phone only", "New phone number", "Open Create account | Enter name, phone, password; leave email empty | Submit", "Account created, signed in, header shows Account; no email required")
T(P, M, "P0", "Sign up with phone and email", "New phone + email", "Fill all fields | Submit", "Account created; both saved (see Account → Profile)")
T(P, M, "P0", "Phone number formats accepted", "New numbers", "Type or paste 9876543210, 09876543210, +91 98765 43210", "All become the same 10-digit number; letters cannot be typed")
T(P, M, "P1", "Invalid phone rejected", "None", "Try 12345, 1234567890 (starts with 1), 9 digits", "Field/message says enter a valid 10-digit mobile number")
T(P, M, "P1", "Invalid email rejected", "None", "Enter 'abc' as email", "Message 'enter a valid email address or leave it empty'")
T(P, M, "P1", "Password rules", "None", "Try 7 characters | Try 8+", "Under 8 rejected with a message; 8+ accepted; Show/Hide button works")
T(P, M, "P0", "Duplicate phone / email blocked", "Existing account", "Sign up again with the same phone; then same email with a new phone", "Message says an account already exists; asks to sign in")
T(P, M, "P0", "Sign in with phone + password", "Account exists", "Login page → enter phone and password", "Signed in; goes to home (or back to the page you came from)")
T(P, M, "P0", "Sign in with email (Use Email-ID)", "Account with email", "Click Use Email-ID | Enter email and password", "Signed in")
T(P, M, "P0", "Wrong password / unknown phone", "None", "Try each", "Same generic message 'Phone number, email or password is incorrect'; no hint which part is wrong")
T(P, M, "P1", "Login rate limit", "None", "Try 9 wrong passwords quickly", "After 8 tries a 'Too many attempts, wait' message; works again after ~15 minutes")
T(P, M, "P0", "Sign out", "Signed in", "Click Sign out (Account page / Admin top bar)", "Signed out everywhere; back button does not show account data; cart badge resets")
T(P, M, "P1", "Session lasts", "Signed in", "Close and reopen the browser the next day (within 7 days)", "Still signed in")
T(P, M, "P1", "After login, return to the previous page", "Signed out, item in cart", "Cart → Continue to checkout → Create account / Log in", "After signing in you land on checkout, not on the home page")
T(P, M, "P1", "Forgot password (email)", "Account with email", "Forgot password → enter email", "Neutral message shown regardless; a reset link is created/sent per email setup")
T(P, M, "P1", "Forgot password for phone-only account", "Phone-only account", "Use Forgot password", "Page explains to contact support; no error")
T(P, M, "P1", "Reset password link", "Valid reset link", "Open link | Set a new password | Log in", "Works once; reused/expired link is rejected")
T(P, M, "P0", "Customer cannot open admin", "Signed in as customer", "Open /admin and /admin/orders", "Redirected/blocked; no admin data visible")
T(P, M, "P0", "Signed-out user cannot open account pages", "Signed out", "Open /account, /account/orders", "Redirected to login")
T(P, M, "P0", "Admin login", "Admin credentials", "Login page → Use Email-ID → admin email + password", "Lands in the Admin dashboard")
T(P, M, "P2", "Disabled user cannot sign in", "Admin disables a customer (Users)", "Try to log in as them", "Sign-in refused")

# ---------------------------------------------------------------- 8 checkout & payment
M = "08 Checkout & payment"; P = "CHK"
T(P, M, "P0", "Sign-up is mandatory before checkout", "Signed out, item in cart", "Cart → Continue to checkout", "Sent to Create account (with a Log in link); guest cannot reach the payment form")
T(P, M, "P0", "Checkout form for new address", "Signed in, item in cart", "Fill name, phone, address, city, state, pincode", "Accepted; phone accepts 10 digits or +91; pincode must be 6 digits")
T(P, M, "P1", "Checkout validation", "None", "Submit with empty/invalid fields", "Readable field errors; entered values are NOT wiped after an error")
T(P, M, "P1", "Saved address", "Account with a saved address", "Open checkout", "Saved address selectable and used; 'new address' also works")
T(P, M, "P2", "Different billing address", "None", "Tick different billing address and fill it", "Order shows both addresses")
T(P, M, "P0", "Razorpay test payment success", "Razorpay TEST keys; card 4111 1111 1111 1111, any future date, CVV 123", "Continue to secure payment | Pay | Wait", "Redirects to order confirmation 'Order confirmed'; cart empty; badge gone; order in Account → Orders as paid")
T(P, M, "P0", "Payment failure keeps the cart", "Razorpay test", "Start payment | Use the failure test card / click Payment failed", "Error shown, order not confirmed, cart items still there, badge unchanged")
T(P, M, "P0", "Close the payment window", "Razorpay test", "Start payment | Close the popup without paying", "Message 'Payment was not completed'; cart still full; can start again")
T(P, M, "P0", "Retry after failure", "After the previous test", "Go to cart → checkout again → pay successfully", "Only one confirmed order; the earlier unpaid order is cancelled and its stock returned")
T(P, M, "P1", "Double click on Pay", "None", "Click Continue to secure payment twice quickly", "Only one order/payment created")
T(P, M, "P0", "UPI / net banking test", "Razorpay test mode", "Pay using the test UPI (success@razorpay) and net banking test bank", "Success confirmed; order paid")
T(P, M, "P0", "Totals match Razorpay amount", "None", "Compare the amount in the Razorpay popup with the cart total", "Exactly equal (in rupees)")
T(P, M, "P0", "Coupon carries into checkout", "Coupon applied in cart", "Continue to checkout and pay", "Discount included in the paid amount and shown on the order page; coupon usage count +1 after payment")
T(P, M, "P1", "Shipping / GST on order", "Two states (home state vs other)", "Place orders to both", "Home state shows CGST+SGST; other state shows IGST; delivery charge per shipping rules")
T(P, M, "P1", "Stock reduces on order", "Product with stock N", "Pay for 1 unit", "Stock becomes N-1; unpaid abandoned orders release stock (after new checkout or 15 min)")
T(P, M, "P1", "Last unit race", "Product with stock 1, two customers", "Both add it | Both pay at the same time", "Only one succeeds; the other sees an out-of-stock message; no negative stock")
T(P, M, "P1", "Reservation expiry", "Abandoned unpaid order", "Wait 15+ minutes (cron runs) | Check stock", "Stock returned; order cancelled; customer sees 'cancelled' on its page")
T(P, M, "P1", "Gateway unavailable message", "Disable Razorpay in settings", "Open checkout", "Razorpay shown as unavailable; button says credentials required; no crash")
T(P, M, "P2", "Snapmint / EMI option", "Not yet active", "Look at payment methods", "Listed as unavailable (until Snapmint is enabled via Razorpay)")
T(P, M, "P1", "Browser back after paying", "Completed payment", "Press Back from the order page", "No duplicate order; cart still empty")
T(P, M, "P0", "Payment confirmed even if the page is closed", "Razorpay test", "Pay | Close the tab right after paying", "Webhook confirms: order becomes CONFIRMED/paid within a minute (check Admin → Orders)")

# ---------------------------------------------------------------- 9 orders & account
M = "09 Order confirmation & My account"; P = "ORD"
T(P, M, "P0", "Order confirmation page", "After a paid order", "Read the page", "Shows order number, items with images, amounts, GST split, address, payment status, tracker, next steps")
T(P, M, "P1", "Confirmation page actions", "Same", "Click Continue shopping, Track order, Print", "Each works; print layout is clean")
T(P, M, "P1", "Receipt link privacy", "Order link", "Open the link signed out / in another account", "Receipt viewable by link only; the account 'Track order' button appears only for the owner")
T(P, M, "P1", "Unpaid order page", "Unpaid order", "Open its confirmation link", "Shows 'awaiting payment' state and a retry option; not 'confirmed'")
T(P, M, "P0", "My orders list and details", "Customer with orders", "Account → Orders → open one", "Only their own orders; statuses and items correct")
T(P, M, "P0", "Cannot see another customer's order", "Two accounts", "Copy an order URL from account A, open as B", "Not found / blocked")
T(P, M, "P1", "Order cancelled by admin", "A paid or unpaid order", "Admin → Orders → open it → set status CANCELLED", "Status changes; stock returns; customer sees it cancelled; a paid order is refunded through the refund flow")
T(P, M, "P1", "Returns and refunds follow the policy", "Delivered order", "Customer contacts support for a return | Admin marks RETURN_REQUESTED → RETURNED and processes the refund", "Status flow allowed; refund amount right; customer notified; outside the return window the team refuses per the policy")
T(P, M, "P1", "Profile update", "Signed in", "Account → Profile → change name/phone", "Saved; phone must be valid and not used by someone else")
T(P, M, "P1", "Change password", "Signed in", "Profile → change password with wrong then right current password", "Wrong current password refused; right one works; new password signs in")
T(P, M, "P1", "Address book", "Signed in", "Add, edit, set default, delete an address", "All work and appear at checkout")
T(P, M, "P1", "Wishlist", "Signed in", "Add/remove items | Add a wishlist item to the cart", "Works and persists")
T(P, M, "P2", "My reviews / notifications", "Signed in", "Open Account → Reviews and Notifications", "Pages load; order updates appear as notifications")

# ---------------------------------------------------------------- 10 admin
M = "10 Admin panel"; P = "ADM"
T(P, M, "P0", "Admin layout", "Admin signed in", "Open several admin pages", "Only the admin top bar and sidebar show (no store header); the active menu item is highlighted; works on phone width", "Desktop + Phone")
T(P, M, "P0", "Admin sign out", "Admin signed in", "Click Sign out in the top bar", "Signed out and sent to the login page")
T(P, M, "P1", "Dashboard", "Some orders exist", "Open Admin home", "Sales, orders, customers, low stock and quick actions load with correct numbers")
T(P, M, "P0", "Create a product", "Admin", "Products → New product → fill name, category, price, stock, photo | Save", "Product appears in Admin list and on the store (if active)")
T(P, M, "P0", "Edit price / stock / status", "Product exists", "Change price, stock, set to draft | Save", "Changes show on the store; draft products hidden from customers")
T(P, M, "P1", "Variants", "Product exists", "Add a variant with its own price/stock | Remove one", "Store shows the variants; totals follow the chosen variant")
T(P, M, "P1", "Duplicate / archive product", "Product exists", "Click duplicate, then archive", "Duplicate created; archived product disappears from store but old orders keep it")
T(P, M, "P1", "Product CSV template, import and export", "Admin", "Download the template | Fill 3 products | Import | Export the catalog", "Rows created; bad rows reported with a clear reason; export matches the catalog")
T(P, M, "P1", "Product photos", "Storage configured", "Upload 1 and 5 photos; delete one", "Photos show on the store; needs Supabase storage settings on Render")
T(P, M, "P1", "Categories, brands, collections", "Admin", "Add / rename / delete each", "Store menus update; deleting one with products is handled safely")
T(P, M, "P0", "Start a sale (guided)", "Products exist", "Offers → choose a category and a % → Start sale", "Prices on the store drop and show a badge; Running sales list shows it")
T(P, M, "P0", "End a sale", "Running sale", "Click End sale on it", "Prices return to normal on the store")
T(P, M, "P1", "Sale with dates", "None", "Create a sale that starts tomorrow and ends later", "Not active today; active in its window; ends automatically")
T(P, M, "P0", "Create / pause / delete coupon", "Admin", "Offers → New coupon → fill code and value | Pause | Delete", "Coupon works in the cart when active; refused when paused")
T(P, M, "P1", "Coupon limits", "Coupon with usage limit and per-user rule", "Use it beyond the limit", "Refused with a clear message")
T(P, M, "P1", "Homepage editor", "Admin", "Storefront → change headline, sections, banners", "Home page reflects the change; Save button shows unsaved state and success message")
T(P, M, "P1", "Announcement bar editor", "Admin", "Homepage → edit messages", "Top scrolling bar updates")
T(P, M, "P1", "Policies & pages editor", "Admin", "Policies → open each policy | Edit text | Save | Reset to default", "Policy page updates; business name/address tokens replaced; reset restores default")
T(P, M, "P1", "Settings (support email/phone, store name)", "Admin", "Settings → change and save", "Contact page and footer show the new details")
T(P, M, "P0", "Order list, search and filter", "Orders exist", "Admin → Orders → filter by status and search by number/phone/email", "Correct results; phone-only customers show their phone")
T(P, M, "P0", "Order status flow", "A paid order", "Move CONFIRMED → PACKED → SHIPPED → OUT_FOR_DELIVERY → DELIVERED", "Each step allowed in order; invalid jumps refused; customer sees status; notifications created")
T(P, M, "P1", "Shipment / tracking details", "Order being shipped", "Add courier and tracking number", "Customer sees them on the order")
T(P, M, "P1", "Export orders CSV", "Orders exist", "Click export", "CSV downloads with correct columns; phone/email present")
T(P, M, "P0", "Refund", "Paid order, Razorpay test", "Create a refund request | Process it", "Refund goes to Razorpay; order/payment status update; customer notified; cannot refund twice")
T(P, M, "P1", "Payments page", "Payments exist", "Admin → Payments", "Shows gateway, status, amounts; matches Razorpay dashboard")
T(P, M, "P1", "Inventory", "Admin", "Adjust stock with a reason", "Stock updates; history shows the change")
T(P, M, "P1", "Reviews moderation", "A pending review", "Approve, then hide it", "Approved shows on the product page; hidden disappears")
T(P, M, "P1", "Customers list", "Customers exist", "Open Customers", "Shows name, phone/email, order count")
T(P, M, "P1", "Users & roles", "Super admin", "Users → change a user's role and disable one", "Role takes effect on next request; product manager cannot see orders they aren't allowed to, etc.")
T(P, M, "P1", "Role restrictions", "PRODUCT_MANAGER and ORDER_MANAGER users", "Sign in as each | Try the other's pages", "Each can only use the areas its role allows")
T(P, M, "P2", "Tax and shipping rules", "Admin", "Edit GST rule and a shipping rule", "New orders use them")
T(P, M, "P2", "Analytics", "Data exists", "Open Analytics", "Charts load; numbers plausible")
T(P, M, "P2", "Audit log", "Admin actions done", "Open Audit logs", "Admin logins and changes are recorded with who and when")

# ---------------------------------------------------------------- 11 payments back-end
M = "11 Payments, webhooks & background jobs"; P = "PAY"
T(P, M, "P0", "Webhook success updates the order", "Order pending, Razorpay test", "Make a test payment but close the tab immediately", "Order becomes paid via the webhook; Razorpay dashboard shows the webhook delivered with 200")
T(P, M, "P1", "Webhook with wrong signature", "None", "Send a fake POST to the webhook URL", "Rejected (400); nothing changes")
T(P, M, "P1", "Duplicate webhook", "Paid order", "Resend the same webhook from Razorpay dashboard", "No double processing: one confirmation, stock and coupon counted once")
T(P, M, "P1", "payment.failed webhook", "Failed test payment", "Check the order", "Payment marked failed; cart kept; stock released later")
T(P, M, "P1", "Reconcile job fixes stuck payments", "A paid order stuck as pending", "Run the reconcile cron manually (Test run)", "Order updated to paid after checking with Razorpay")
T(P, M, "P1", "Release reservations job", "Expired unpaid order", "Test run the release job", "Order cancelled, stock returned")
T(P, M, "P1", "Cron endpoints refuse strangers", "None", "POST to each cron URL without the Bearer header", "401 for all three")
T(P, M, "P2", "Notification dispatch", "Pending notifications", "Test run the job", "Notifications marked sent; no errors")

# ---------------------------------------------------------------- 12 legal
M = "12 Policies, cookies & legal"; P = "LEG"
T(P, M, "P0", "All 5 policy pages open", "None", "Open Privacy, Terms, Shipping, Refund/Returns, Cookie policy from the footer", "Each loads; text reads correctly; no {placeholders} visible; brand name is Mobile & More")
T(P, M, "P0", "Business details filled in", "Admin", "Policies → set legal business name, address, GST number, contact", "Policy pages and contact page show them; no defaults like the admin's personal email")
T(P, M, "P1", "Cookie banner first visit", "New/private window", "Open the site", "Banner appears with Accept all / Essential only and a Cookie Policy link; does not block browsing")
T(P, M, "P1", "Cookie choice remembered", "Banner shown", "Click Accept all | Reload | Open another page", "Banner does not return; 'Cookie settings' in the footer reopens it")
T(P, M, "P1", "Essential only", "New window", "Choose Essential only", "Banner closes; cart and login still work")
T(P, M, "P1", "Contact page", "None", "Open Contact us", "Shows real email/phone/address from Settings; links (mailto:, tel:) work")
T(P, M, "P2", "Consent text at sign-up / checkout", "None", "Look under Create account and Continue to payment", "Links to Terms and Privacy Policy present and open")
T(P, M, "P1", "Legal review", "Owner/lawyer", "Read all 5 policies once", "Wording approved (returns window, warranty, shipping times match the real business)")

# ---------------------------------------------------------------- 13 security
M = "13 Security & privacy"; P = "SEC"
T(P, M, "P0", "Admin API blocked for customers", "Customer signed in", "Call /api/admin/products (browser or DevTools)", "403 Administrator access is required")
T(P, M, "P0", "Signed-out API access", "Signed out", "Call /api/account/profile", "401 Please sign in")
T(P, M, "P0", "Cross-site request blocked", "None", "From another website/DevTools send a POST to /api/cart/items with a different Origin", "403 Cross-site request rejected")
T(P, M, "P1", "Session cookie flags", "Signed in", "DevTools → Application → Cookies", "'session' is HttpOnly, Secure, SameSite=Lax")
T(P, M, "P1", "Tampered session refused", "Signed in", "Edit a character in the session cookie | Reload", "Treated as signed out")
T(P, M, "P1", "Script injection in text fields", "None", "Use <img src=x onerror=alert(1)> as name, address and review text", "Shown as plain text; no popup anywhere, including Admin")
T(P, M, "P1", "Password hashes never exposed", "None", "Check API responses for users/orders", "No passwordHash/IMEI/tokens in any response")
T(P, M, "P1", "Secrets not in the browser", "None", "View page source and network responses", "No API keys, DATABASE_URL, CRON_SECRET anywhere")
T(P, M, "P1", "Rate limits", "None", "Repeat sign-up, login, forgot-password rapidly", "Limits kick in with a clear message")
T(P, M, "P2", "Security headers", "None", "Check response headers (securityheaders.com)", "X-Frame-Options, Referrer-Policy, Permissions-Policy present")
T(P, M, "P1", "Razorpay uses TEST vs LIVE keys correctly", "Before launch", "Check keys on Render", "Test keys while testing; live keys only after KYC; webhook secret matches")
T(P, M, "P1", "Personal data handling", "Owner", "Check who can see customer phone numbers/addresses", "Only admin roles; support process documented")

# ---------------------------------------------------------------- 14 responsive / browsers
M = "14 Devices, browsers & layout"; P = "DEV"
T(P, M, "P0", "Phone layout (Android Chrome)", "Android phone", "Walk Home → Product → Cart → Sign up → Checkout", "No horizontal scroll, no overlapped or cut text, buttons easy to tap", "Android Chrome")
T(P, M, "P0", "iPhone layout (Safari)", "iPhone", "Same walk-through", "Same; iOS keyboard does not hide fields; sticky header ok", "iPhone Safari")
T(P, M, "P1", "Tablet layout", "iPad/Android tablet", "Browse home, listing, product", "2-3 column grids, header intact", "Tablet")
T(P, M, "P0", "Mac Chrome and Safari", "Mac", "Home, cart, login, policy pages", "Footer at the bottom; density right; fonts load", "Mac Chrome + Safari")
T(P, M, "P1", "Windows Edge and Firefox", "Windows", "Same walk-through", "No layout differences that hurt use", "Edge + Firefox")
T(P, M, "P1", "Different zoom levels", "Desktop", "Set browser zoom to 80%, 100%, 125%", "Site stays usable at each; no overlap", "Desktop")
T(P, M, "P1", "Large monitor (1920+)", "Big screen", "Open Home and listing", "Content centred/scaled well; not stretched or tiny", "1920x1080")
T(P, M, "P2", "Dark mode / OS settings", "OS dark mode", "Open the site", "Still readable (no white-on-white text)")
T(P, M, "P2", "Slow network", "Chrome DevTools Slow 3G", "Open Home and a product", "Skeleton/loading state shows; nothing breaks; images load progressively")
T(P, M, "P2", "Keyboard use", "Desktop", "Tab through header, forms, checkout", "Every control reachable; focus outline visible; Enter submits forms")
T(P, M, "P2", "Screen reader basics", "VoiceOver/NVDA", "Open sign-in", "Fields have labels; errors announced")
T(P, M, "P2", "Print order receipt", "Order page", "Click Print", "Clean printable layout without header/buttons")

# ---------------------------------------------------------------- 15 seo/perf
M = "15 SEO & performance"; P = "SEO"
T(P, M, "P1", "Page titles and descriptions", "None", "Check home, listing, product, policy pages", "Unique titles including 'Mobile & More'; sensible descriptions")
T(P, M, "P1", "sitemap.xml and robots.txt", "None", "Open /sitemap.xml and /robots.txt", "Sitemap lists products/categories/policies; robots allows the store and blocks /admin, /api")
T(P, M, "P1", "Render address not indexed", "None", "Open the Render address", "Redirects to the main site (no duplicate site in Google)")
T(P, M, "P2", "Lighthouse (mobile)", "Chrome", "Run Lighthouse on Home and a product", "Performance 70+, Accessibility 90+, no console errors")
T(P, M, "P2", "Favicon and social preview", "None", "Check the browser tab icon and share the link in WhatsApp", "Icon and title/image show correctly")
T(P, M, "P2", "Broken links crawl", "Tool or manual", "Click through the header, footer and category links", "No 404s")

# ---------------------------------------------------------------- 16 go-live
M = "16 Go-live checklist"; P = "LIV"
T(P, M, "P0", "Remove demo products", "Real catalog ready", "Run the remove-demo script (npm run db:remove-demo) after adding real products", "Only real products remain; homepage offers use real items")
T(P, M, "P0", "Real prices, stock, photos", "Owner", "Review every product in Admin", "Prices/GST/warranty/condition correct; every product has photos")
T(P, M, "P0", "Razorpay LIVE mode", "KYC approved", "Switch keys and webhook secret to live on Render; place a small real order and refund it", "Real payment succeeds; webhook received; refund works")
T(P, M, "P0", "Admin password changed and secrets files deleted", "Owner", "Change the admin password | Delete the secrets .txt files from Documents", "New password works; old one refused; no secrets left on the PC")
T(P, M, "P1", "Support email/phone/address real", "Owner", "Settings and Policies pages", "Real, monitored contact details everywhere")
T(P, M, "P1", "Upgrade or keep-alive for Render", "Owner", "Decide: paid Render plan or keep-alive job", "No cold-start delays for customers")
T(P, M, "P1", "Storage for product photos", "Owner", "Set SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY, STORAGE_BUCKET on Render", "Admin photo upload works")
T(P, M, "P1", "Domain auto-renewal", "Hostinger", "Turn on auto-renew (expires 2027-03-30)", "Enabled")
T(P, M, "P1", "Final full regression", "QA team", "Run every P0 and P1 case above on phone and desktop", "All pass; no open Blocker/High bugs")

HEADER = ["ID", "Module", "Priority", "Test case", "Preconditions", "Steps", "Expected result", "Device / browser", "Status", "Tester", "Date tested", "Bug / Trello link", "Notes"]
out = os.path.join(os.path.dirname(__file__), "manual-test-cases.csv")
with open(out, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(HEADER); w.writerows(rows)
print(len(rows), "test cases ->", out)
for prefix, count in counters.items(): print(prefix, count)


# ---------------------------------------------------------------- workbook: one tab per module + a live Summary tab
try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    raise SystemExit("CSV written. Install openpyxl (pip install openpyxl) to also build the workbook.")

TAB = {"01": "01 Deployment", "02": "02 Header & search", "03": "03 Home page", "04": "04 Listing & filters", "05": "05 Product page", "06": "06 Cart",
       "07": "07 Sign up & sign in", "08": "08 Checkout & payment", "09": "09 Orders & account", "10": "10 Admin panel", "11": "11 Payments & jobs",
       "12": "12 Policies & cookies", "13": "13 Security", "14": "14 Devices & browsers", "15": "15 SEO & performance", "16": "16 Go-live"}
WIDTHS = [10, 26, 9, 34, 30, 52, 52, 18, 11, 14, 13, 24, 30]
DARK = PatternFill("solid", fgColor="14211F")
wb = Workbook()
summary = wb.active; summary.title = "Summary"
modules = []
for row in rows:
    if row[1] not in modules: modules.append(row[1])
for module in modules:
    ws = wb.create_sheet(TAB[module[:2]])
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = DARK; cell.alignment = Alignment(vertical="center", wrap_text=True)
    mine = [r for r in rows if r[1] == module]
    for r in mine: ws.append(r)
    for i, width in enumerate(WIDTHS, 1): ws.column_dimensions[get_column_letter(i)].width = width
    for line in ws.iter_rows(min_row=2):
        for cell in line: cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "E2"; ws.auto_filter.ref = ws.dimensions
    last = len(mine) + 1
    dv = DataValidation(type="list", formula1='"Not run,Pass,Fail,Blocked"', allow_blank=False); ws.add_data_validation(dv); dv.add(f"I2:I{last + 100}")
    for text, color in (("Pass", "C6EFCE"), ("Fail", "FFC7CE"), ("Blocked", "FFEB9C")):
        ws.conditional_formatting.add(f"I2:I{last + 100}", CellIsRule(operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
    for text, color in (("P0", "FFC7CE"), ("P1", "FFEB9C"), ("P2", "DDEBF7")):
        ws.conditional_formatting.add(f"C2:C{last}", CellIsRule(operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))

summary.append(["Mobile & More - manual test progress (updates itself as testers set Status on each tab)"]); summary["A1"].font = Font(bold=True, size=14)
summary.append([])
summary.append(["Module tab", "Total", "Pass", "Fail", "Blocked", "Not run", "% done"])
for cell in summary[3]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = DARK
first = 4
for module in modules:
    name = TAB[module[:2]]; n = summary.max_row + 1
    ref = f"'{name}'!I:I"
    summary.append([name, f'=COUNTA(\'{name}\'!A:A)-1', f'=COUNTIF({ref},"Pass")', f'=COUNTIF({ref},"Fail")', f'=COUNTIF({ref},"Blocked")', f'=COUNTIF({ref},"Not run")', f"=IF(B{n}=0,0,(C{n}+D{n}+E{n})/B{n})"])
    summary[f"A{n}"].hyperlink = f"#'{name}'!A1"
end = summary.max_row; n = end + 1
summary.append(["ALL MODULES"] + [f"=SUM({c}{first}:{c}{end})" for c in "BCDEF"] + [f"=IF(B{n}=0,0,(C{n}+D{n}+E{n})/B{n})"])
for cell in summary[n]: cell.font = Font(bold=True)
for r in range(first, n + 1): summary[f"G{r}"].number_format = "0%"
summary.column_dimensions["A"].width = 30
for c in "BCDEFG": summary.column_dimensions[c].width = 11
summary.append([]); summary.append(["Priority guide: P0 = must pass before launch, P1 = important, P2 = nice to have. Log every failed case as a Trello card titled with its Test ID."])
xlsx = os.path.join(os.path.dirname(__file__), "manual-test-cases.xlsx")
wb.save(xlsx); print("workbook ->", xlsx, "tabs:", wb.sheetnames)
